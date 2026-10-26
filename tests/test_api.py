import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from app.config import settings

def test_api_health(client: TestClient):
    response = client.get("/api/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert data["database_connected"] is True
    assert "version" in data

def test_api_upload_success(client: TestClient):
    csv_data = "customer_id,name,email,created_at\nCUST-001,Alice Smith,alice@example.com,2024-01-01\n"
    files = {"file": ("customers.csv", io.BytesIO(csv_data.encode("utf-8")), "text/csv")}
    response = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)

    assert response.status_code == 201
    data = response.json()
    assert data["status"] == "completed"
    assert data["total_rows"] == 1
    assert data["accepted_rows"] == 1
    assert data["rejected_rows"] == 0
    assert data["is_repeated_upload"] is False

def test_api_upload_unsupported_file_extension(client: TestClient):
    files = {"file": ("data.txt", io.BytesIO(b"some,data\n"), "text/plain")}
    response = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    assert response.status_code == 415
    assert "only standard csv" in response.json()["detail"].lower()

def test_api_upload_empty_file(client: TestClient):
    files = {"file": ("empty.csv", io.BytesIO(b""), "text/csv")}
    response = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()

def test_api_upload_invalid_dataset(client: TestClient):
    files = {"file": ("test.csv", io.BytesIO(b"col1,col2\n1,2"), "text/csv")}
    response = client.post("/api/uploads", data={"dataset_type": "invalid_dataset"}, files=files)
    assert response.status_code == 422
    assert "invalid dataset_type" in response.json()["detail"].lower()

def test_api_jobs_list_and_detail(client: TestClient):
    # Upload one job first
    csv_data = "customer_id,name,email,created_at\nCUST-010,Dan,dan@example.com,2024-01-01\n"
    files = {"file": ("cust_job.csv", io.BytesIO(csv_data.encode("utf-8")), "text/csv")}
    upload_res = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    job_id = upload_res.json()["job_id"]

    # List jobs
    list_res = client.get("/api/jobs")
    assert list_res.status_code == 200
    jobs_data = list_res.json()
    assert jobs_data["total"] >= 1
    assert any(j["id"] == job_id for j in jobs_data["items"])

    # Get single job details
    detail_res = client.get(f"/api/jobs/{job_id}")
    assert detail_res.status_code == 200
    detail = detail_res.json()
    assert detail["id"] == job_id
    assert detail["filename"] == "cust_job.csv"

def test_api_job_not_found(client: TestClient):
    response = client.get("/api/jobs/999999")
    assert response.status_code == 404
    assert "was not found" in response.json()["detail"].lower()

def test_api_validation_errors_endpoints(client: TestClient):
    # Upload invalid customer row
    bad_csv = "customer_id,name,email,created_at\nCUST-011,Eve,not-an-email,2024-01-01\n"
    files = {"file": ("bad_email.csv", io.BytesIO(bad_csv.encode("utf-8")), "text/csv")}
    upload_res = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    job_id = upload_res.json()["job_id"]

    # Job specific errors
    job_errs_res = client.get(f"/api/jobs/{job_id}/errors")
    assert job_errs_res.status_code == 200
    data = job_errs_res.json()
    assert data["total"] >= 1
    assert data["items"][0]["error_code"] == "INVALID_EMAIL_FORMAT"

    # Global errors list
    all_errs_res = client.get("/api/errors?error_code=INVALID_EMAIL_FORMAT")
    assert all_errs_res.status_code == 200
    assert all_errs_res.json()["total"] >= 1

def test_api_clean_records(client: TestClient):
    csv_data = "customer_id,name,email,created_at\nCUST-050,Frank Castle,frank@example.com,2024-01-01\n"
    files = {"file": ("clean_cust.csv", io.BytesIO(csv_data.encode("utf-8")), "text/csv")}
    client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)

    # Fetch clean records
    rec_res = client.get("/api/records?dataset_type=customers&q=Castle")
    assert rec_res.status_code == 200
    rec_data = rec_res.json()
    assert rec_data["total"] == 1
    assert rec_data["items"][0]["customer_id"] == "CUST-050"

def test_api_clean_records_unsupported_dataset(client: TestClient):
    response = client.get("/api/records?dataset_type=nonexistent")
    assert response.status_code == 422

def test_api_dashboard_endpoints(client: TestClient):
    csv_data = "sale_id,customer_id,amount,sale_date,currency\nSALE-001,CUST-001,250.00,2024-02-01,USD\n"
    files = {"file": ("sales_dash.csv", io.BytesIO(csv_data.encode("utf-8")), "text/csv")}
    client.post("/api/uploads", data={"dataset_type": "sales"}, files=files)

    summary_res = client.get("/api/dashboard/summary")
    assert summary_res.status_code == 200
    summary = summary_res.json()
    assert summary["kpis"]["total_records_processed"] >= 1
    assert summary["kpis"]["accepted_records"] >= 1

    trends_res = client.get("/api/dashboard/trends")
    assert trends_res.status_code == 200
    trends = trends_res.json()
    assert len(trends["trends"]) >= 1

def test_api_reprocess_endpoint(client: TestClient):
    csv_data = "product_id,product_name,quantity,unit_price,updated_at\nPROD-1,Desk,5,199.99,2024-03-01\n"
    files = {"file": ("inventory.csv", io.BytesIO(csv_data.encode("utf-8")), "text/csv")}
    upload_res = client.post("/api/uploads", data={"dataset_type": "inventory"}, files=files)
    job_id = upload_res.json()["job_id"]

    reprocess_res = client.post(f"/api/jobs/{job_id}/reprocess")
    assert reprocess_res.status_code == 200
    data = reprocess_res.json()
    assert data["id"] == job_id
    assert data["status"] == "completed"
    assert data["accepted_rows"] == 1
    assert data["duplicate_rows"] == 0

    # Reprocessed jobs retain their accepted-row export.
    clean_res = client.get(f"/api/jobs/{job_id}/download/clean")
    assert clean_res.status_code == 200
    assert b"PROD-1" in clean_res.content

def test_download_clean_records_customers(client: TestClient):
    csv_data = "customer_id,name,email,phone,created_at\nCUST-DL-1,Grace Hopper,grace@example.com,+123456789,2024-01-15\n"
    files = {"file": ("customers_dl.csv", io.BytesIO(csv_data.encode("utf-8")), "text/csv")}
    upload_res = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    assert upload_res.status_code == 201
    job_id = upload_res.json()["job_id"]

    res = client.get(f"/api/jobs/{job_id}/download/clean")
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]
    assert f'filename="clean_customers_job_{job_id}.csv"' in res.headers["content-disposition"]
    text = res.text
    assert "customer_id,name,email,phone,created_at" in text
    assert "CUST-DL-1" in text
    assert "grace@example.com" in text

def test_download_clean_records_sales_and_inventory(client: TestClient):
    # Test sales download
    sales_csv = "sale_id,customer_id,amount,sale_date,currency\nSALE-DL-1,CUST-DL-1,149.50,2024-02-10,USD\n"
    files = {"file": ("sales_dl.csv", io.BytesIO(sales_csv.encode("utf-8")), "text/csv")}
    up_sales = client.post("/api/uploads", data={"dataset_type": "sales"}, files=files)
    assert up_sales.status_code == 201
    sales_job_id = up_sales.json()["job_id"]

    res_sales = client.get(f"/api/jobs/{sales_job_id}/download/clean")
    assert res_sales.status_code == 200
    assert "sale_id,customer_id,amount,sale_date,currency" in res_sales.text
    assert "SALE-DL-1" in res_sales.text

    # Test inventory download
    inv_csv = "product_id,product_name,quantity,unit_price,updated_at\nPROD-DL-1,Office Chair,25,89.99,2024-03-01\n"
    files = {"file": ("inv_dl.csv", io.BytesIO(inv_csv.encode("utf-8")), "text/csv")}
    up_inv = client.post("/api/uploads", data={"dataset_type": "inventory"}, files=files)
    assert up_inv.status_code == 201
    inv_job_id = up_inv.json()["job_id"]

    res_inv = client.get(f"/api/jobs/{inv_job_id}/download/clean")
    assert res_inv.status_code == 200
    assert "product_id,product_name,quantity,unit_price,updated_at" in res_inv.text
    assert "PROD-DL-1" in res_inv.text

def test_download_clean_records_zero_accepted_returns_400(client: TestClient):
    bad_csv = "customer_id,name,email,created_at\nCUST-BAD-1,Bad Guy,invalid-email,2024-01-01\n"
    files = {"file": ("bad_only.csv", io.BytesIO(bad_csv.encode("utf-8")), "text/csv")}
    upload_res = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    assert upload_res.status_code == 201
    job_id = upload_res.json()["job_id"]
    assert upload_res.json()["accepted_rows"] == 0

    res = client.get(f"/api/jobs/{job_id}/download/clean")
    assert res.status_code == 400
    assert "no clean records" in res.json()["detail"].lower()

def test_download_validation_errors_csv(client: TestClient):
    bad_csv = "customer_id,name,email,created_at\nCUST-ERR-1,Error User,invalid-format,2024-01-01\n"
    files = {"file": ("errors_dl.csv", io.BytesIO(bad_csv.encode("utf-8")), "text/csv")}
    upload_res = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    assert upload_res.status_code == 201
    job_id = upload_res.json()["job_id"]

    res = client.get(f"/api/jobs/{job_id}/download/errors")
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]
    assert f'filename="validation_errors_job_{job_id}.csv"' in res.headers["content-disposition"]
    text = res.text
    assert "job_id,row_number,dataset_type,field_name,error_code,error_message,original_value" in text
    assert "INVALID_EMAIL_FORMAT" in text
    assert "invalid-format" in text

def test_download_validation_errors_zero_errors_returns_400(client: TestClient):
    clean_csv = "customer_id,name,email,created_at\nCUST-PERF-1,Perfect Person,perfect@example.com,2024-01-01\n"
    files = {"file": ("perfect.csv", io.BytesIO(clean_csv.encode("utf-8")), "text/csv")}
    upload_res = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    assert upload_res.status_code == 201
    job_id = upload_res.json()["job_id"]
    assert upload_res.json()["rejected_rows"] == 0
    assert upload_res.json()["duplicate_rows"] == 0

    res = client.get(f"/api/jobs/{job_id}/download/errors")
    assert res.status_code == 400
    assert "no validation errors" in res.json()["detail"].lower()

def test_download_job_summary_excel(client: TestClient):
    mixed_csv = (
        "customer_id,name,email,created_at\n"
        "CUST-MIX-1,Valid One,valid1@example.com,2024-01-01\n"
        "CUST-MIX-2,Bad Email,invalid-email-address,2024-01-01\n"
    )
    files = {"file": ("mixed_summary.csv", io.BytesIO(mixed_csv.encode("utf-8")), "text/csv")}
    upload_res = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    assert upload_res.status_code == 201
    job_id = upload_res.json()["job_id"]

    res = client.get(f"/api/jobs/{job_id}/download/summary")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f'filename="job_{job_id}_report.xlsx"' in res.headers["content-disposition"]

    # Verify genuine Excel workbook by parsing with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    sheet_names = wb.sheetnames
    assert sheet_names == ["Job Summary", "Clean Records", "Validation Errors"]

    # 1. Check Job Summary sheet
    ws_summary = wb["Job Summary"]
    summary_text = [str(cell.value) for row in ws_summary.iter_rows() for cell in row if cell.value is not None]
    assert any(f"#{job_id}" in t for t in summary_text)
    assert any("mixed_summary.csv" in t for t in summary_text)
    assert any("Customers" in t for t in summary_text)
    assert any("PASSED" in t for t in summary_text)

    # 2. Check Clean Records sheet
    ws_clean = wb["Clean Records"]
    assert ws_clean.freeze_panes == "A2"
    clean_headers = [cell.value for cell in ws_clean[1]]
    assert clean_headers == ["customer_id", "name", "email", "phone", "created_at"]
    # Row 2 contains the 1 accepted customer
    assert ws_clean.cell(row=2, column=1).value == "CUST-MIX-1"

    # 3. Check Validation Errors sheet
    ws_errors = wb["Validation Errors"]
    assert ws_errors.freeze_panes == "A2"
    error_headers = [cell.value for cell in ws_errors[1]]
    assert error_headers == ["job_id", "row_number", "dataset_type", "field_name", "error_code", "error_message", "original_value"]
    # Row 2 contains the invalid email error
    assert ws_errors.cell(row=2, column=5).value == "INVALID_EMAIL_FORMAT"
    assert ws_errors.cell(row=2, column=7).value == "invalid-email-address"


def test_download_job_summary_excel_empty_clean_and_empty_errors(client: TestClient):
    # Job with 0 accepted records (all rejected)
    bad_csv = "customer_id,name,email,created_at\nCUST-BAD,Bad Person,bad-email,2024-01-01\n"
    files = {"file": ("all_bad.csv", io.BytesIO(bad_csv.encode("utf-8")), "text/csv")}
    upload_bad = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    bad_job_id = upload_bad.json()["job_id"]

    res_bad = client.get(f"/api/jobs/{bad_job_id}/download/summary")
    assert res_bad.status_code == 200
    wb_bad = openpyxl.load_workbook(io.BytesIO(res_bad.content))
    ws_clean_bad = wb_bad["Clean Records"]
    # Headers are present
    assert [cell.value for cell in ws_clean_bad[1]] == ["customer_id", "name", "email", "phone", "created_at"]
    # Informative message in row 2
    assert "no clean records were accepted" in str(ws_clean_bad["A2"].value).lower()

    # Job with 0 errors (all clean)
    clean_csv = "customer_id,name,email,created_at\nCUST-GOOD,Good Person,good@example.com,2024-01-01\n"
    files = {"file": ("all_good.csv", io.BytesIO(clean_csv.encode("utf-8")), "text/csv")}
    upload_good = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    good_job_id = upload_good.json()["job_id"]

    res_good = client.get(f"/api/jobs/{good_job_id}/download/summary")
    assert res_good.status_code == 200
    wb_good = openpyxl.load_workbook(io.BytesIO(res_good.content))
    ws_errors_good = wb_good["Validation Errors"]
    # Headers are present
    assert [cell.value for cell in ws_errors_good[1]] == ["job_id", "row_number", "dataset_type", "field_name", "error_code", "error_message", "original_value"]
    # Informative message in row 2
    assert "clean run" in str(ws_errors_good["A2"].value).lower()


def test_download_job_summary_excel_sales_and_inventory(client: TestClient):
    # Sales job
    sales_csv = "sale_id,customer_id,amount,sale_date,currency\nSALE-EX-1,CUST-001,120.50,2024-02-01,USD\n"
    files = {"file": ("sales_ex.csv", io.BytesIO(sales_csv.encode("utf-8")), "text/csv")}
    up_sales = client.post("/api/uploads", data={"dataset_type": "sales"}, files=files)
    sales_id = up_sales.json()["job_id"]

    res_sales = client.get(f"/api/jobs/{sales_id}/download/summary")
    assert res_sales.status_code == 200
    wb_sales = openpyxl.load_workbook(io.BytesIO(res_sales.content))
    ws_clean_sales = wb_sales["Clean Records"]
    assert [cell.value for cell in ws_clean_sales[1]] == ["sale_id", "customer_id", "amount", "sale_date", "currency"]
    assert ws_clean_sales.cell(row=2, column=1).value == "SALE-EX-1"
    assert ws_clean_sales.cell(row=2, column=3).value == 120.50

    # Inventory job
    inv_csv = "product_id,product_name,quantity,unit_price,updated_at\nPROD-EX-1,Keyboard,15,45.00,2024-03-01\n"
    files = {"file": ("inv_ex.csv", io.BytesIO(inv_csv.encode("utf-8")), "text/csv")}
    up_inv = client.post("/api/uploads", data={"dataset_type": "inventory"}, files=files)
    inv_id = up_inv.json()["job_id"]

    res_inv = client.get(f"/api/jobs/{inv_id}/download/summary")
    assert res_inv.status_code == 200
    wb_inv = openpyxl.load_workbook(io.BytesIO(res_inv.content))
    ws_clean_inv = wb_inv["Clean Records"]
    assert [cell.value for cell in ws_clean_inv[1]] == ["product_id", "product_name", "quantity", "unit_price", "updated_at"]
    assert ws_clean_inv.cell(row=2, column=1).value == "PROD-EX-1"
    assert ws_clean_inv.cell(row=2, column=3).value == 15
    assert ws_clean_inv.cell(row=2, column=4).value == 45.00


def test_download_job_summary_excel_isolation(client: TestClient):
    # Job 1
    csv_1 = "customer_id,name,email,created_at\nCUST-EX-A,Alice Excel,alice.excel@example.com,2024-01-01\n"
    files_1 = {"file": ("job_1_ex.csv", io.BytesIO(csv_1.encode("utf-8")), "text/csv")}
    res_1 = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files_1)
    job_1_id = res_1.json()["job_id"]

    # Job 2
    csv_2 = "customer_id,name,email,created_at\nCUST-EX-B,Bob Excel,bob.excel@example.com,2024-01-01\n"
    files_2 = {"file": ("job_2_ex.csv", io.BytesIO(csv_2.encode("utf-8")), "text/csv")}
    res_2 = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files_2)
    job_2_id = res_2.json()["job_id"]

    # Check Job 1 workbook Clean Records
    res_wb1 = client.get(f"/api/jobs/{job_1_id}/download/summary")
    wb1 = openpyxl.load_workbook(io.BytesIO(res_wb1.content))
    ws1_clean = wb1["Clean Records"]
    ws1_vals = [cell.value for cell in ws1_clean["A"]]
    assert "CUST-EX-A" in ws1_vals
    assert "CUST-EX-B" not in ws1_vals

    # Check Job 2 workbook Clean Records
    res_wb2 = client.get(f"/api/jobs/{job_2_id}/download/summary")
    wb2 = openpyxl.load_workbook(io.BytesIO(res_wb2.content))
    ws2_clean = wb2["Clean Records"]
    ws2_vals = [cell.value for cell in ws2_clean["A"]]
    assert "CUST-EX-B" in ws2_vals
    assert "CUST-EX-A" not in ws2_vals

def test_download_nonexistent_job_returns_404(client: TestClient):
    res_clean = client.get("/api/jobs/999999/download/clean")
    assert res_clean.status_code == 404
    assert "not found" in res_clean.json()["detail"].lower()

    res_err = client.get("/api/jobs/999999/download/errors")
    assert res_err.status_code == 404
    assert "not found" in res_err.json()["detail"].lower()

    res_sum = client.get("/api/jobs/999999/download/summary")
    assert res_sum.status_code == 404
    assert "not found" in res_sum.json()["detail"].lower()

def test_download_job_isolation(client: TestClient):
    # Job A
    csv_a = "customer_id,name,email,created_at\nCUST-ISO-A,Alice Isolation,alice.iso@example.com,2024-01-01\n"
    files_a = {"file": ("job_a.csv", io.BytesIO(csv_a.encode("utf-8")), "text/csv")}
    res_a = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files_a)
    job_a_id = res_a.json()["job_id"]

    # Job B
    csv_b = "customer_id,name,email,created_at\nCUST-ISO-B,Bob Isolation,bob.iso@example.com,2024-01-01\n"
    files_b = {"file": ("job_b.csv", io.BytesIO(csv_b.encode("utf-8")), "text/csv")}
    res_b = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files_b)
    job_b_id = res_b.json()["job_id"]

    # Download clean records for Job A
    download_a = client.get(f"/api/jobs/{job_a_id}/download/clean")
    assert download_a.status_code == 200
    assert "CUST-ISO-A" in download_a.text
    assert "CUST-ISO-B" not in download_a.text

    # Download clean records for Job B
    download_b = client.get(f"/api/jobs/{job_b_id}/download/clean")
    assert download_b.status_code == 200
    assert "CUST-ISO-B" in download_b.text
    assert "CUST-ISO-A" not in download_b.text

def test_download_csv_escaping(client: TestClient):
    # Customer name with quotes, commas, and special punctuation
    csv_data = 'customer_id,name,email,created_at\nCUST-ESC-1,"O\'Connor, Jr. ""The Great""",oconnor@example.com,2024-01-01\n'
    files = {"file": ("special_chars.csv", io.BytesIO(csv_data.encode("utf-8")), "text/csv")}
    upload_res = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    assert upload_res.status_code == 201
    job_id = upload_res.json()["job_id"]

    res = client.get(f"/api/jobs/{job_id}/download/clean")
    assert res.status_code == 200
    
    # Parse through standard csv.reader to verify correct escaping
    import csv
    reader = list(csv.reader(io.StringIO(res.text)))
    assert len(reader) == 2  # header + 1 row
    header = reader[0]
    row = reader[1]
    name_idx = header.index("name")
    assert row[name_idx] == 'O\'Connor, Jr. "The Great"'



def test_upload_rejects_oversized_file_before_full_read(client: TestClient, monkeypatch):
    """Reject a payload above the configured limit without buffering the full file."""
    monkeypatch.setattr(settings, "UPLOAD_MAX_SIZE_BYTES", 16)
    csv_data = b"customer_id,name,email,created_at\nCUST-1,Alice,alice@example.com,2024-01-01\n"
    files = {"file": ("large.csv", io.BytesIO(csv_data), "text/csv")}
    response = client.post("/api/uploads", data={"dataset_type": "customers"}, files=files)
    assert response.status_code == 413
    assert client.get("/api/jobs").json()["total"] == 0


@pytest.mark.parametrize("url", [
    "/api/jobs?start_date=not-a-date",
    "/api/records?dataset_type=customers&start_date=not-a-date",
    "/api/errors?start_date=not-a-date",
    "/api/dashboard/summary?start_date=not-a-date",
    "/api/dashboard/trends?end_date=2026-02-30",
])
def test_invalid_date_filters_return_422(client: TestClient, url: str):
    response = client.get(url)
    assert response.status_code == 422
