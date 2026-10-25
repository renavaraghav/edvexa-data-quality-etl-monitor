import csv
import io
import json
import logging
from typing import Any
from fastapi import HTTPException, status
from sqlalchemy.orm import Session
from sqlalchemy import select

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import settings
from app.models import ProcessingJob, JobCleanRecord, ValidationError
from app.services.validation import DATASET_REQUIRED_COLUMNS

logger = logging.getLogger("edvexa.export")

DATASET_EXPORT_COLUMNS: dict[str, list[str]] = {
    "customers": ["customer_id", "name", "email", "phone", "created_at"],
    "sales": ["sale_id", "customer_id", "amount", "sale_date", "currency"],
    "inventory": ["product_id", "product_name", "quantity", "unit_price", "updated_at"],
}

def backfill_job_clean_records_if_needed(job: ProcessingJob, db: Session) -> list[JobCleanRecord]:
    """
    Safe backwards-compatibility resolver:
    If a completed job has accepted_rows > 0 but was processed before JobCleanRecord was introduced,
    re-extract the accepted rows from its securely saved source artifact.
    """
    stmt = (
        select(JobCleanRecord)
        .where(JobCleanRecord.job_id == job.id)
        .order_by(JobCleanRecord.row_number.asc())
    )
    records = list(db.scalars(stmt).all())
    if records:
        return records

    # If job has no accepted rows or no artifact, return empty
    if job.accepted_rows == 0:
        return []

    artifact_name = f"job_{job.id}_{job.file_hash[:12]}.csv"
    artifact_path = settings.UPLOADS_DIR / artifact_name
    if not artifact_path.exists():
        return []

    try:
        from app.services.csv_import import parse_and_validate_csv
        from app.services.validation import validate_row_by_dataset, DATASET_PRIMARY_KEYS
        from app.models import utc_now

        with open(artifact_path, "rb") as f:
            content = f.read()

        extraction = parse_and_validate_csv(content, job.filename, job.dataset_type)
        if not extraction.is_valid:
            return []

        primary_key_field = DATASET_PRIMARY_KEYS.get(job.dataset_type, "")
        seen_keys: set[str] = set()
        
        # Check against existing error rows for this job
        error_rows = {
            e.row_number for e in db.scalars(
                select(ValidationError).where(ValidationError.job_id == job.id)
            ).all()
        }

        created_records = []
        for parsed_row in extraction.rows:
            if parsed_row.line_number in error_rows:
                continue

            val_res = validate_row_by_dataset(job.dataset_type, parsed_row.data)
            if not val_res.is_valid or not val_res.cleaned_data:
                continue

            key_val = val_res.cleaned_data.get(primary_key_field)
            if not key_val or key_val in seen_keys:
                continue

            seen_keys.add(key_val)
            rec = JobCleanRecord(
                job_id=job.id,
                row_number=parsed_row.line_number,
                dataset_type=job.dataset_type,
                business_key=str(key_val),
                record_json=json.dumps(val_res.cleaned_data, default=str),
                created_at=utc_now()
            )
            created_records.append(rec)
            db.add(rec)

        if created_records:
            db.commit()
            return created_records

    except Exception as e:
        logger.warning(f"Could not backfill clean records for Job #{job.id}: {e}")

    return []

def sanitize_cell_value(val: Any) -> Any:
    """
    Mitigate CSV / Formula Injection (CWE-1236).
    If a string value starts with =, +, -, @, \\t, or \\r, prefix with single quote
    to neutralize formula execution in spreadsheet applications.
    """
    if isinstance(val, str) and len(val) > 0 and val[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + val
    return val

def generate_clean_records_csv(job: ProcessingJob, db: Session) -> tuple[str, str]:
    """
    Generate downloadable CSV containing only clean, normalized records accepted in this job.
    Properly escapes commas, quotes, and newlines.
    """
    if job.accepted_rows == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"No clean records to download. Job #{job.id} has 0 accepted records."
        )

    records = backfill_job_clean_records_if_needed(job, db)
    if not records:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"No clean records found for Job #{job.id}."
        )

    columns = DATASET_EXPORT_COLUMNS.get(job.dataset_type, ["data"])

    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(columns)

    for rec in records:
        try:
            row_dict = json.loads(rec.record_json)
        except Exception:
            row_dict = {}
        writer.writerow([sanitize_cell_value(row_dict.get(col, "")) for col in columns])

    filename = f"clean_{job.dataset_type}_job_{job.id}.csv"
    return output.getvalue(), filename

def generate_validation_errors_csv(job: ProcessingJob, db: Session) -> tuple[str, str]:
    """
    Generate downloadable CSV containing all validation and duplicate detection errors for this job.
    Includes every recorded error across all rows.
    """
    errors = job.validation_errors
    if not errors:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"No validation errors to download. Job #{job.id} contains 0 errors."
        )

    columns = [
        "job_id",
        "row_number",
        "dataset_type",
        "field_name",
        "error_code",
        "error_message",
        "original_value",
    ]

    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(columns)

    for err in errors:
        writer.writerow([
            err.job_id,
            err.row_number,
            job.dataset_type,
            err.field_name,
            err.error_code,
            sanitize_cell_value(err.error_message),
            sanitize_cell_value(err.original_value if err.original_value is not None else ""),
        ])

    filename = f"validation_errors_job_{job.id}.csv"
    return output.getvalue(), filename

def generate_job_summary_json(job: ProcessingJob, db: Session) -> tuple[dict[str, Any], str]:
    """
    Generate downloadable JSON document summarizing job execution metrics and metadata.
    """
    total = job.total_rows
    acc = job.accepted_rows
    rej = job.rejected_rows
    dup = job.duplicate_rows
    skp = job.skipped_rows

    assert total == (acc + rej + dup + skp), (
        f"Accounting mismatch for Job #{job.id}: total={total} vs sum={acc + rej + dup + skp}"
    )

    errors_by_field: dict[str, int] = {}
    errors_by_code: dict[str, int] = {}
    for err in job.validation_errors:
        f_name = err.field_name or "unknown"
        e_code = err.error_code or "UNKNOWN"
        errors_by_field[f_name] = errors_by_field.get(f_name, 0) + 1
        errors_by_code[e_code] = errors_by_code.get(e_code, 0) + 1

    pass_rate = round((acc / total) * 100, 2) if total > 0 else 100.0

    summary = {
        "job_id": job.id,
        "filename": job.filename,
        "dataset_type": job.dataset_type,
        "status": job.status,
        "sha256_checksum": job.file_hash,
        "timestamps": {
            "started_at": job.started_at.isoformat() if job.started_at else None,
            "completed_at": job.completed_at.isoformat() if job.completed_at else None,
            "duration_seconds": job.processing_duration,
        },
        "metrics": {
            "total_rows": total,
            "accepted_rows": acc,
            "rejected_rows": rej,
            "duplicate_rows": dup,
            "skipped_rows": skp,
            "data_quality_rate_pct": pass_rate,
        },
        "reconciliation": {
            "reconciled": total == (acc + rej + dup + skp),
            "formula": "total_rows == accepted_rows + rejected_rows + duplicate_rows + skipped_rows",
            "total_rows": total,
            "sum_of_parts": acc + rej + dup + skp,
            "difference": total - (acc + rej + dup + skp),
        },
        "errors_summary": {
            "total_validation_errors": len(job.validation_errors),
            "by_field": errors_by_field,
            "by_code": errors_by_code,
        },
    }

    filename = f"job_{job.id}_summary.json"
    return summary, filename


def generate_job_excel_report(job: ProcessingJob, db: Session) -> tuple[bytes, str]:
    """
    Generate a comprehensive, professionally styled Microsoft Excel (.xlsx) workbook containing:
    - Sheet 1: Job Summary & Accounting Reconciliation
    - Sheet 2: Clean Records (job-scoped accepted records)
    - Sheet 3: Validation Errors (all validation & duplicate detection failures)
    """
    wb = Workbook()

    # Styling definitions
    navy_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    section_banner_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    sub_banner_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    subtle_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    banner_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    sub_banner_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_label_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=10, color="1E293B")
    italic_empty_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    pass_font = Font(name="Calibri", size=10, bold=True, color="166534")
    fail_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    grid_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ---------------------------------------------------------
    # SHEET 1: Job Summary
    # ---------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Job Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_summary.merge_cells("A1:B1")
    title_cell = ws_summary["A1"]
    title_cell.value = "EDVEXA DATA QUALITY & ETL MONITOR - JOB AUDIT REPORT"
    title_cell.font = banner_font
    title_cell.fill = section_banner_fill
    title_cell.alignment = align_center
    ws_summary.row_dimensions[1].height = 28

    # Section 1: Execution Metadata
    ws_summary.merge_cells("A3:B3")
    meta_banner = ws_summary["A3"]
    meta_banner.value = "1. JOB EXECUTION METADATA"
    meta_banner.font = sub_banner_font
    meta_banner.fill = sub_banner_fill
    meta_banner.alignment = align_left
    ws_summary.row_dimensions[3].height = 20

    started_str = job.started_at.strftime("%Y-%m-%d %H:%M:%S UTC") if job.started_at else "N/A"
    completed_str = job.completed_at.strftime("%Y-%m-%d %H:%M:%S UTC") if job.completed_at else "N/A"
    duration_str = f"{job.processing_duration:.3f} s" if job.processing_duration is not None else "N/A"

    metadata_rows = [
        ("Job ID", f"#{job.id}"),
        ("Original Filename", job.filename),
        ("Dataset Type", job.dataset_type.title()),
        ("Processing Status", job.status.upper()),
        ("SHA-256 Checksum", job.file_hash),
        ("Start Timestamp", started_str),
        ("Completion Timestamp", completed_str),
        ("Processing Duration", duration_str),
    ]

    curr_row = 4
    for label, val in metadata_rows:
        cell_lbl = ws_summary.cell(row=curr_row, column=1, value=label)
        cell_lbl.font = bold_label_font
        cell_lbl.border = grid_border
        cell_lbl.alignment = align_left
        cell_lbl.fill = subtle_row_fill

        cell_val = ws_summary.cell(row=curr_row, column=2, value=val)
        cell_val.font = regular_font
        cell_val.border = grid_border
        cell_val.alignment = align_left
        ws_summary.row_dimensions[curr_row].height = 18
        curr_row += 1

    # Section 2: Pipeline Metrics
    curr_row += 1
    ws_summary.merge_cells(f"A{curr_row}:B{curr_row}")
    metrics_banner = ws_summary[f"A{curr_row}"]
    metrics_banner.value = "2. PIPELINE EXECUTION METRICS"
    metrics_banner.font = sub_banner_font
    metrics_banner.fill = sub_banner_fill
    metrics_banner.alignment = align_left
    ws_summary.row_dimensions[curr_row].height = 20
    curr_row += 1

    total = job.total_rows
    acc = job.accepted_rows
    rej = job.rejected_rows
    dup = job.duplicate_rows
    skp = job.skipped_rows
    err_count = len(job.validation_errors)
    pass_rate_str = f"{(acc / total * 100):.2f}%" if total > 0 else "100.00%"

    metrics_data = [
        ("Total Rows Processed", total),
        ("Accepted Clean Records", acc),
        ("Validation Rejections", rej),
        ("Duplicate Records Prevented", dup),
        ("Skipped Rows", skp),
        ("Validation Error Count", err_count),
        ("Data Quality Pass Rate", pass_rate_str),
    ]

    for label, val in metrics_data:
        cell_lbl = ws_summary.cell(row=curr_row, column=1, value=label)
        cell_lbl.font = bold_label_font
        cell_lbl.border = grid_border
        cell_lbl.alignment = align_left
        cell_lbl.fill = subtle_row_fill

        cell_val = ws_summary.cell(row=curr_row, column=2, value=val)
        cell_val.font = regular_font
        cell_val.border = grid_border
        if isinstance(val, (int, float)):
            cell_val.alignment = align_right
            cell_val.number_format = "#,##0"
        else:
            cell_val.alignment = align_left
        ws_summary.row_dimensions[curr_row].height = 18
        curr_row += 1

    # Section 3: Accounting Reconciliation
    curr_row += 1
    ws_summary.merge_cells(f"A{curr_row}:B{curr_row}")
    recon_banner = ws_summary[f"A{curr_row}"]
    recon_banner.value = "3. ACCOUNTING RECONCILIATION AUDIT"
    recon_banner.font = sub_banner_font
    recon_banner.fill = sub_banner_fill
    recon_banner.alignment = align_left
    ws_summary.row_dimensions[curr_row].height = 20
    curr_row += 1

    sum_parts = acc + rej + dup + skp
    diff = total - sum_parts
    recon_passed = (diff == 0)

    recon_data = [
        ("Reconciliation Invariant", "total_rows = accepted + rejected + duplicates + skipped"),
        ("Total Rows Ingested", total),
        ("Sum of Mutually Exclusive Parts", sum_parts),
        ("Reconciliation Difference", diff),
        ("Reconciliation Status", "PASSED - Exact 100% Row Accounting" if recon_passed else "FAILED - Row Count Mismatch"),
    ]

    for label, val in recon_data:
        cell_lbl = ws_summary.cell(row=curr_row, column=1, value=label)
        cell_lbl.font = bold_label_font
        cell_lbl.border = grid_border
        cell_lbl.alignment = align_left
        cell_lbl.fill = subtle_row_fill

        cell_val = ws_summary.cell(row=curr_row, column=2, value=val)
        cell_val.border = grid_border
        if label == "Reconciliation Status":
            cell_val.font = pass_font if recon_passed else fail_font
            cell_val.fill = pass_fill if recon_passed else fail_fill
            cell_val.alignment = align_center
        elif isinstance(val, (int, float)):
            cell_val.font = regular_font
            cell_val.alignment = align_right
            cell_val.number_format = "#,##0"
        else:
            cell_val.font = regular_font
            cell_val.alignment = align_left
        ws_summary.row_dimensions[curr_row].height = 18
        curr_row += 1

    ws_summary.column_dimensions["A"].width = 36
    ws_summary.column_dimensions["B"].width = 54

    # ---------------------------------------------------------
    # SHEET 2: Clean Records
    # ---------------------------------------------------------
    ws_clean = wb.create_sheet(title="Clean Records")
    ws_clean.views.sheetView[0].showGridLines = True
    clean_columns = DATASET_EXPORT_COLUMNS.get(job.dataset_type, ["data"])

    # Header Row
    ws_clean.row_dimensions[1].height = 24
    for col_idx, col_name in enumerate(clean_columns, start=1):
        c = ws_clean.cell(row=1, column=col_idx, value=col_name)
        c.font = header_font
        c.fill = navy_header_fill
        c.alignment = align_center
        c.border = grid_border

    clean_records = backfill_job_clean_records_if_needed(job, db)
    if clean_records:
        for row_idx, rec in enumerate(clean_records, start=2):
            ws_clean.row_dimensions[row_idx].height = 19
            try:
                row_dict = json.loads(rec.record_json)
            except Exception:
                row_dict = {}

            for col_idx, col_name in enumerate(clean_columns, start=1):
                val = row_dict.get(col_name, "")
                cell = ws_clean.cell(row=row_idx, column=col_idx)
                cell.border = grid_border
                cell.font = regular_font

                # Formats based on field type
                if col_name in ("amount", "unit_price"):
                    try:
                        cell.value = float(val) if val != "" else None
                        cell.number_format = "#,##0.00"
                        cell.alignment = align_right
                    except (ValueError, TypeError):
                        cell.value = sanitize_cell_value(val)
                        cell.alignment = align_right
                elif col_name == "quantity":
                    try:
                        cell.value = int(val) if val != "" else None
                        cell.number_format = "#,##0"
                        cell.alignment = align_right
                    except (ValueError, TypeError):
                        cell.value = sanitize_cell_value(val)
                        cell.alignment = align_right
                else:
                    cell.value = sanitize_cell_value(val)
                    cell.alignment = align_left

        max_col_letter = get_column_letter(len(clean_columns))
        ws_clean.auto_filter.ref = f"A1:{max_col_letter}{len(clean_records) + 1}"
    else:
        # Empty state with informative message
        ws_clean.row_dimensions[2].height = 22
        ws_clean.merge_cells(f"A2:{get_column_letter(len(clean_columns))}2")
        empty_cell = ws_clean["A2"]
        empty_cell.value = "No clean records were accepted for this job (0 rows)."
        empty_cell.font = italic_empty_font
        empty_cell.alignment = align_center
        empty_cell.fill = subtle_row_fill
        empty_cell.border = grid_border

    ws_clean.freeze_panes = "A2"

    # Auto-adjust column widths
    for col in ws_clean.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.value is not None:
                val_str = str(cell.value)
                line_len = max(len(line) for line in val_str.split("\n"))
                max_len = max(max_len, line_len)
        ws_clean.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 14)

    # ---------------------------------------------------------
    # SHEET 3: Validation Errors
    # ---------------------------------------------------------
    ws_errors = wb.create_sheet(title="Validation Errors")
    ws_errors.views.sheetView[0].showGridLines = True
    error_columns = [
        "job_id",
        "row_number",
        "dataset_type",
        "field_name",
        "error_code",
        "error_message",
        "original_value",
    ]

    ws_errors.row_dimensions[1].height = 24
    for col_idx, col_name in enumerate(error_columns, start=1):
        c = ws_errors.cell(row=1, column=col_idx, value=col_name)
        c.font = header_font
        c.fill = navy_header_fill
        c.alignment = align_center
        c.border = grid_border

    errors = job.validation_errors
    if errors:
        for row_idx, err in enumerate(errors, start=2):
            ws_errors.row_dimensions[row_idx].height = 20
            row_vals = [
                err.job_id,
                err.row_number,
                job.dataset_type,
                err.field_name,
                err.error_code,
                err.error_message,
                err.original_value if err.original_value is not None else "",
            ]

            for col_idx, val in enumerate(row_vals, start=1):
                clean_val = sanitize_cell_value(val) if isinstance(val, str) else val
                cell = ws_errors.cell(row=row_idx, column=col_idx, value=clean_val)
                cell.border = grid_border
                cell.font = regular_font

                if col_idx in (1, 2):  # job_id, row_number
                    cell.alignment = align_center
                    cell.number_format = "#,##0"
                elif col_idx == 6:  # error_message
                    cell.alignment = align_wrap
                elif col_idx == 5:  # error_code
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        max_err_col_letter = get_column_letter(len(error_columns))
        ws_errors.auto_filter.ref = f"A1:{max_err_col_letter}{len(errors) + 1}"
    else:
        # Empty state with informative message
        ws_errors.row_dimensions[2].height = 22
        ws_errors.merge_cells(f"A2:{get_column_letter(len(error_columns))}2")
        empty_err_cell = ws_errors["A2"]
        empty_err_cell.value = "Clean run: No validation errors recorded for this job (0 errors)."
        empty_err_cell.font = italic_empty_font
        empty_err_cell.alignment = align_center
        empty_err_cell.fill = subtle_row_fill
        empty_err_cell.border = grid_border

    ws_errors.freeze_panes = "A2"

    # Auto-adjust column widths for Errors sheet
    for col in ws_errors.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.value is not None:
                val_str = str(cell.value)
                line_len = max(len(line) for line in val_str.split("\n"))
                max_len = max(max_len, line_len)
        ws_errors.column_dimensions[col_letter].width = max(min(max_len + 4, 55), 14)

    # Save to in-memory bytes
    output = io.BytesIO()
    wb.save(output)
    excel_bytes = output.getvalue()

    filename = f"job_{job.id}_report.xlsx"
    return excel_bytes, filename
